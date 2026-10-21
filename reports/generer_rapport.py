# -*- coding: utf-8 -*-
"""Generateur de rapport EcoDrive robuste aux accents/encodages des cles JSON.
Placez ce fichier dans reports/ sous le nom generer_rapport.py puis lancez :
    py -3.12 reports/generer_rapport.py
"""
from pathlib import Path
import json, unicodedata
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "reports" / "resultats.json"
OUT = ROOT / "reports" / "Rapport_EcoDrive.docx"


def norm(s):
    s = str(s).replace("²", "2").replace("€", "eur")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.lower().replace("_", " ").split())


def get(d, *names, default=None):
    if not isinstance(d, dict):
        return default
    nd = {norm(k): v for k, v in d.items()}
    for name in names:
        if norm(name) in nd:
            return nd[norm(name)]
    return default


def num(x, dec=2):
    if x is None:
        return "-"
    try:
        return f"{float(x):.{dec}f}".replace(".", ",")
    except Exception:
        return str(x)


def pct(x, dec=0):
    try:
        return num(float(x) * 100, dec) + " %"
    except Exception:
        return "-"


data = json.loads(RESULTS.read_text(encoding="utf-8"))
D = get(data, "donnees", default={})
G = get(data, "regression", default={})
X = get(data, "decision", default={})
comparaison = get(G, "comparaison", default=[]) or []
roi = get(X, "roi", default=[]) or []
styles = get(X, "styles", default=[]) or []

def scenario(prefix):
    for r in roi:
        if str(get(r, "scenario", default="")).startswith(prefix):
            return r
    return {}

central = scenario("Central")

# ---------------- document ----------------
doc = Document()
styles_doc = doc.styles
styles_doc["Normal"].font.name = "Calibri"
styles_doc["Normal"].font.size = Pt(10.5)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("EcoDrive")
r.bold = True; r.font.size = Pt(28)
p = doc.add_paragraph("Prédire la consommation de carburant d'une flotte pour détecter les dérives et comparer les styles de conduite")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("Enzo BASLY · Guillaume · Rayan")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Résumé", level=1)
mt = get(G, "test_modele", default={}) or {}
mr = get(G, "test_regle", default={}) or {}
doc.add_paragraph(
    f"Le modèle retenu prévoit la consommation d'un trajet avec une MAE de {num(get(mt,'MAE'))} L/100 km, "
    f"contre {num(get(mr,'MAE'))} L/100 km pour la règle actuelle. "
    f"Le seuil d'alerte est fixé à +{num(float(get(X,'seuil',default=0))*100,0)} %, avec un recall de {pct(get(X,'recall',default=0))}."
)
doc.add_paragraph(
    f"Dans le scénario central, le gain annuel estimé est de {num(get(central,'gain annuel (eur)','gain annuel (€)'),0)} € "
    f"pour un coût de première année de {num(get(central,'cout an 1 (eur)','coût an 1 (€)'),0)} €, "
    f"soit un ROI sur 3 ans de {num(get(central,'ROI 3 ans (%)'),0)} %."
)

doc.add_heading("1. Problème métier", level=1)
doc.add_paragraph("EcoDrive vise une flotte de véhicules de service. Le gestionnaire veut repérer les véhicules qui consomment anormalement et identifier les conduites qui augmentent la consommation.")
doc.add_paragraph("Question étudiée : quelle consommation est normale pour un trajet donné ? La cible est continue (L/100 km), donc le problème principal est une régression supervisée.")

doc.add_heading("2. Données et préparation", level=1)
doc.add_paragraph(
    f"Après préparation, le jeu contient {get(D,'n_trajets',default='-')} trajets et {get(D,'n_vehicules',default='-')} véhicules. "
    "Les mesures OBD sont regroupées par trajet. Les trajets trop courts ou incohérents sont retirés."
)
doc.add_paragraph("Les variables qui servent directement à calculer la cible, comme le débit carburant ou le débit d'air, ne sont pas utilisées comme variables explicatives afin d'éviter une fuite de données.")

doc.add_heading("3. Modélisation", level=1)
doc.add_paragraph("Le découpage est chronologique : apprentissage sur les trajets les plus anciens, test sur les plus récents. Les transformations sont réalisées dans un pipeline, puis plusieurs modèles sont comparés par validation croisée temporelle.")
if comparaison:
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, h in enumerate(["Modèle", "MAE train", "MAE CV", "Écart-type", "R² CV"]):
        table.rows[0].cells[i].text = h
    for c in comparaison:
        row = table.add_row().cells
        vals = [
            get(c, "modele", "modèle", default="-"),
            num(get(c, "MAE train")),
            num(get(c, "MAE CV")),
            num(get(c, "ecart-type", "écart-type")),
            num(get(c, "R2 CV", "R² CV")),
        ]
        for i, v in enumerate(vals): row[i].text = str(v)

doc.add_paragraph(f"Modèle retenu : {get(G,'modele_retenu','modèle retenu',default='-')}.")
doc.add_paragraph(f"Sur le test : MAE = {num(get(mt,'MAE'))} L/100 km, RMSE = {num(get(mt,'RMSE'))}, R² = {num(get(mt,'R2','R²'))}.")

doc.add_heading("4. Alertes de surconsommation", level=1)
doc.add_paragraph(
    f"Le seuil est choisi sur le train puis vérifié sur le test. Seuil retenu : +{num(float(get(X,'seuil',default=0))*100,0)} %. "
    f"Recall test : {pct(get(X,'recall',default=0))}. Precision test : {pct(get(X,'precision',default=0))}. "
    f"Fausses alertes estimées : {num(get(X,'FP/an','fp_an'),0)} par an."
)
doc.add_paragraph(
    f"Coût annuel estimé du système : {num(get(X,'cout_modele','coût modèle'),0)} €, "
    f"contre {num(get(X,'cout_sans','coût sans'),0)} € sans système selon les hypothèses retenues."
)

doc.add_heading("5. Styles de conduite", level=1)
if styles:
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, h in enumerate(["Contexte", "Calme", "Normal", "Nerveux", "Écart nerveux/calme"]):
        table.rows[0].cells[i].text = h
    for s in styles:
        row = table.add_row().cells
        vals = [
            get(s,"contexte",default="-"), num(get(s,"calme")), num(get(s,"normal")), num(get(s,"nerveux")),
            pct(get(s,"ecart nerveux / calme","écart nerveux / calme",default=0))
        ]
        for i,v in enumerate(vals): row[i].text = str(v)
doc.add_paragraph("La comparaison montre qu'à type de trajet égal, une conduite plus nerveuse est associée à une consommation plus élevée. Le modèle est ensuite utilisé pour estimer un potentiel d'économie à véhicule et trajet comparables.")

doc.add_heading("6. ROI", level=1)
if roi:
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i,h in enumerate(["Scénario", "Gain annuel", "Coût an 1", "ROI an 1", "ROI 3 ans"]):
        table.rows[0].cells[i].text = h
    for s in roi:
        row = table.add_row().cells
        vals = [
            get(s,"scenario",default="-"),
            num(get(s,"gain annuel (eur)","gain annuel (€)"),0) + " €",
            num(get(s,"cout an 1 (eur)","coût an 1 (€)"),0) + " €",
            num(get(s,"ROI an 1 (%)"),0) + " %",
            num(get(s,"ROI 3 ans (%)"),0) + " %",
        ]
        for i,v in enumerate(vals): row[i].text = str(v)

doc.add_heading("7. Limites et recommandation", level=1)
for txt in [
    "Les défauts sont simulés car le dataset ne contient pas d'étiquettes de panne.",
    "Les données VED proviennent des États-Unis ; un pilote sur une flotte française est nécessaire.",
    "Le nombre de fausses alertes doit être réduit avant un déploiement général.",
    "Recommandation : pilote de 3 mois sur 20 véhicules, mesure avant/après, puis décision de déploiement."
]:
    doc.add_paragraph(txt, style="List Bullet")

doc.add_heading("Sources et IA générative", level=1)
doc.add_paragraph("Source des données : Vehicle Energy Dataset (VED), G. S. Oh, D. J. LeBlanc, H. Peng, 2020.")
doc.add_paragraph("Microsoft Copilot a été utilisé pour aider à structurer le code et la rédaction. Les analyses sont exécutées dans les notebooks et doivent être comprises par l'équipe.")

doc.save(OUT)
print("OK rapport ->", OUT)
